#!/usr/bin/env python3
"""
Import pipeline workbook -> data dictionary (menambal temuan J audit).

Langkah 1-2 dari "Prompt Claude" di GDD: AUDIT semua sheet (header, satuan, jumlah
baris, merged cells) dan bangun DATA DICTIONARY — TANPA menebak header atau mengubah
dosis. Baris multi-bahan yang ambigu ditandai untuk review, tidak dipaksa diparse.

Pemakaian:
  1. Letakkan Std_Produksi_Pupuk_Herbisida_v3.xlsx di folder ini (atau beri path).
  2. pip3 install openpyxl
  3. python3 tools/import_workbook.py [path_ke_xlsx]

Output: tools/data_dictionary.json  (ringkasan tiap sheet untuk ditinjau manusia)

Konversi ke data/*.json runtime BARU dilakukan setelah dictionary ditinjau &
pemetaan kolom dikonfirmasi (lihat bagian MAPPING di bawah).
"""
import json, os, sys, glob

HERE = os.path.dirname(__file__)
DEFAULT_NAMES = ["Std_Produksi_Pupuk_Herbisida_v3.xlsx", "*.xlsx"]

# Sheet yang dipakai game (lihat GDD v0.3 bagian 2). Kolom TIDAK diasumsikan di sini.
EXPECTED_SHEETS = [
    "Parameter", "Std Produksi", "T_Produksi", "Std Pupuk", "T_Pupuk",
    "Std Semprot", "T_Semprot", "Mkns_Blanket", "Detil merk", "Detil merk Mkns",
    "Harga Herbisida", "T_Harga", "Bibitan", "Ulat Api", "Hama Lain", "Catatan Perubahan",
]

def find_workbook(arg):
    if arg and os.path.exists(arg):
        return arg
    for pat in DEFAULT_NAMES:
        hits = glob.glob(os.path.join(HERE, pat)) + glob.glob(os.path.join(HERE, "..", "..", pat))
        if hits:
            return hits[0]
    return None

def main():
    arg = sys.argv[1] if len(sys.argv) > 1 else None
    wb_path = find_workbook(arg)
    if not wb_path:
        print("❌ Workbook .xlsx belum ditemukan.")
        print("   Letakkan Std_Produksi_Pupuk_Herbisida_v3.xlsx di tools/ lalu jalankan lagi.")
        print("   (Temuan J audit: file ini belum dilampirkan ke proyek.)")
        return 1
    try:
        import openpyxl
    except ImportError:
        print("❌ Perlu openpyxl. Jalankan: pip3 install openpyxl")
        return 1

    print("📖 Membaca:", wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)   # data_only=False -> lihat rumus
    wb_vals = openpyxl.load_workbook(wb_path, data_only=True)  # hasil rumus

    dictionary = {"source_filename": os.path.basename(wb_path), "sheets": {}}
    for name in wb.sheetnames:
        ws = wb[name]
        ws_v = wb_vals[name]
        header = [c.value for c in next(ws.iter_rows(max_row=1), [])]
        sample = []
        for row in ws_v.iter_rows(min_row=2, max_row=6, values_only=True):
            sample.append(list(row))
        dictionary["sheets"][name] = {
            "expected_by_game": name in EXPECTED_SHEETS,
            "dimensions": ws.dimensions,
            "n_rows": ws.max_row,
            "n_cols": ws.max_column,
            "merged_cells": [str(m) for m in ws.merged_cells.ranges],
            "header_row1": header,
            "sample_rows_2_6": sample,
            "has_formulas": any(
                isinstance(c.value, str) and c.value.startswith("=")
                for r in ws.iter_rows() for c in r
            ),
            "review_flags": _flags(name, ws),
        }

    missing = [s for s in EXPECTED_SHEETS if s not in wb.sheetnames]
    if missing:
        dictionary["MISSING_EXPECTED_SHEETS"] = missing

    out = os.path.join(HERE, "data_dictionary.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(dictionary, f, ensure_ascii=False, indent=2, default=str)
    print("✅ Data dictionary ditulis:", out)
    print("   Tinjau dulu (terutama merged cells & baris multi-bahan Std Semprot)")
    print("   sebelum konversi ke data/*.json. JANGAN menebak header/dosis.")
    return 0

def _flags(name, ws):
    flags = []
    if ws.merged_cells.ranges:
        flags.append("ADA merged cells — cek baris multi-bahan (GDD 10).")
    if name in ("Std Semprot", "T_Semprot"):
        flags.append("Normalisasi baris multi-bahan: satu rekomendasi bisa >1 komponen.")
    if name in ("Ulat Api", "Hama Lain"):
        flags.append("Satuan campuran (per ha / konsentrasi / umpan per pokok) — validasi wajib.")
    if name in ("Harga Herbisida", "T_Harga"):
        flags.append("Pisahkan harga aktual vs estimasi + tahun harga.")
    return flags

# ---------------------------------------------------------------------------
# MAPPING (diisi SETELAH dictionary ditinjau). Template per entitas GDD §4:
#   production_standard <- Std Produksi / T_Produksi
#   fertilizer_standard <- Std Pupuk / T_Pupuk
#   control_standard    <- Std Semprot / Ulat Api / Hama Lain
#   product_catalog     <- Detil merk / Harga Herbisida
# Setiap baris output WAJIB simpan: source_sheet, source_row, source_version,
# unit asli, dan validation_status="needs_validation".
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    sys.exit(main())
