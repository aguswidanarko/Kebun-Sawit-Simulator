#!/usr/bin/env python3
"""Membuat Excel daftar item yang HARUS direview secara agronomi."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(__file__)
DATA = os.path.join(HERE, "..", "data")
OUT = os.path.join(HERE, "..", "..", "Review_Agronomi_Kebun_Sawit_Simulator.xlsx")

def load(n): return json.load(open(os.path.join(DATA, n), encoding="utf-8"))

GREEN = "2E7D32"; LIGHT = "E8F5E9"; YELLOW = "FFF9C4"; GREY = "ECEFF1"
FONT = "Arial"
thin = Side(style="thin", color="B0BEC5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name=FONT, bold=True, size=13, color=GREEN)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

def make_sheet(name, headers, rows, widths, note="", review_cols=True):
    ws = wb.create_sheet(name)
    full = list(headers)
    if review_cols:
        full += ["Status Review\n(OK/Perbaiki/Tolak)", "Nilai Terverifikasi", "Sumber / Label / Registrasi", "Catatan Reviewer"]
    title(ws, name.replace("_", " ") + ("  —  " + note if note else ""), len(full))
    hr = 2
    for i, h in enumerate(full, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header(ws, hr, len(full))
    for r, row in enumerate(rows, hr + 1):
        for cidx in range(len(full)):
            val = row[cidx] if cidx < len(row) else ""
            cell = ws.cell(row=r, column=cidx + 1, value=val)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if cidx >= len(headers):
                cell.fill = PatternFill("solid", fgColor=YELLOW)
    for i in range(len(full)):
        ws.column_dimensions[get_column_letter(i + 1)].width = widths[i] if i < len(widths) else 20
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(full))}{hr}"
    return ws

# ------------------------------------------------- 1) RINGKASAN
ws = wb.active; ws.title = "Ringkasan Review"
title(ws, "DAFTAR REVIEW AGRONOMI — Kebun Sawit Simulator", 7)
intro = ("Semua item di bawah memakai data/nilai yang BELUM final secara agronomi. Tinjau sebelum "
         "game dipakai sebagai materi edukasi publik (lihat GDD v0.3 §7). Prioritas: TINGGI = menyangkut "
         "dosis/legalitas/keselamatan pestisida; SEDANG = angka ekonomi/parameter; RENDAH = penyeimbang gameplay.")
ws.merge_cells("A2:G2"); ws["A2"] = intro
ws["A2"].font = Font(name=FONT, size=10, italic=True); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 48
heads = ["No", "Kategori", "Sheet detail", "Jml item", "Status sekarang", "Yang harus diverifikasi", "Prioritas"]
for i, h in enumerate(heads, 1): ws.cell(row=3, column=i, value=h)
style_header(ws, 3, 7)

seeds = load("seed_catalog.json")
prod = load("product_catalog.json"); prod = [r for r in prod if "trade_name" in r]
ctrl = load("control_standard.json"); ctrl = [r for r in ctrl if "active_ingredient" in r]
pestc = load("pest_control.json"); pestc = [r for r in pestc if r.get("validation_status") == "needs_review"]
pcat = [r for r in load("pest_catalog.json") if "id" in r]

summary = [
    ["Hama/penyakit (data mentah)", "Hama_Penyakit_Mentah", len(pestc), "needs_review (merged cells, satuan campuran)",
     "Dosis, satuan (per ha/pokok/liter), bahan aktif vs nama dagang, status registrasi", "TINGGI"],
    ["Harga & registrasi produk pestisida", "Harga_Produk", len(prod), "estimasi (dari T_Harga)",
     "Harga aktual, status registrasi/izin edar, target yang diizinkan per label", "TINGGI"],
    ["Norma semprot gulma", "Norma_Semprot_Gulma", len(ctrl), "from_workbook (standar perusahaan)",
     "Kesesuaian dosis/konsentrasi dgn label, APD, keselamatan, badan air", "TINGGI"],
    ["Katalog varietas kecambah", "Varietas_Kecambah", len(seeds), "gameplay / estimasi / needs_verification",
     "Harga jual aktual, plafon produksi resmi, rendemen, umur panen, toleransi penyakit", "SEDANG"],
    ["Parameter ekonomi & placeholder", "Parameter_Placeholder", None, "placeholder gameplay",
     "Harga pupuk Rp/kg (tak ada di workbook), harga TBS, biaya obat hama", "SEDANG"],
    ["Katalog hama (angka gameplay)", "Katalog_Hama_Gameplay", len(pcat), "gameplay",
     "Peluang muncul, laju kerusakan, kecocokan metode/bahan dgn kondisi nyata", "RENDAH"],
]
for i, s in enumerate(summary, 1):
    vals = [i] + s
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=3 + i, column=c, value=v)
        cell.font = Font(name=FONT, size=10, bold=(c == 7))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER
        if c == 7:
            color = {"TINGGI": "FFCDD2", "SEDANG": "FFF9C4", "RENDAH": "C8E6C9"}.get(v, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=color)
for i, w in enumerate([5, 30, 24, 9, 30, 42, 11], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A4"

# ------------------------------------------------- 2) HAMA MENTAH
rows = [[r.get("source_sheet"), r.get("source_row"), r.get("diagnosis_group"), r.get("target_organism_raw"),
         r.get("plant_age_raw", ""), r.get("method"), r.get("trade_name"), r.get("active_ingredient"),
         r.get("concentration_raw"), r.get("dose_raw"), r.get("volume_raw", ""), r.get("applicator"),
         r.get("interval_raw"), r.get("notes")] for r in pestc]
make_sheet("Hama_Penyakit_Mentah",
           ["Sheet asal", "Baris asal", "Grup diagnosis", "Organisme (mentah)", "Umur", "Metode",
            "Nama dagang", "Bahan aktif", "Konsentrasi", "Dosis", "Volume", "Aplikator", "Interval/Rotasi", "Keterangan"],
           rows, [11, 8, 16, 26, 10, 14, 22, 22, 12, 12, 10, 16, 16, 30],
           note="PRIORITAS TINGGI — bersihkan satuan & verifikasi dosis/registrasi sebelum dipakai")

# ------------------------------------------------- 3) HARGA PRODUK
rows = [[r.get("active_ingredient"), r.get("trade_name"), r.get("unit"), r.get("price"),
         r.get("price_year"), r.get("price_status"), r.get("price_source"), r.get("source_row")] for r in prod]
make_sheet("Harga_Produk",
           ["Bahan aktif", "Nama dagang", "Satuan", "Harga (Rp)", "Thn harga", "Status harga", "Sumber harga", "Baris asal"],
           rows, [22, 30, 8, 14, 9, 14, 24, 9],
           note="PRIORITAS TINGGI — verifikasi harga aktual + status registrasi/izin edar")

# ------------------------------------------------- 4) NORMA SEMPROT
rows = [[r.get("plant_age_min"), r.get("plant_age_max"), r.get("target_location"), r.get("target_organism"),
         r.get("mix_option"), r.get("active_ingredient"), r.get("dose_value"), r.get("dose_unit"),
         r.get("per_basis"), r.get("water_l_ha_min"), r.get("water_l_ha_max"), r.get("concentration_pct"),
         r.get("equipment"), r.get("source_row")] for r in ctrl]
make_sheet("Norma_Semprot_Gulma",
           ["Umur min", "Umur max", "Lokasi sasaran", "Gulma", "Opsi campuran", "Bahan aktif", "Dosis", "Satuan",
            "Basis", "Vol min (L/ha)", "Vol max (L/ha)", "Konsentrasi %", "Alat", "Baris asal"],
           rows, [9, 9, 22, 22, 11, 20, 9, 8, 8, 12, 12, 12, 18, 9],
           note="PRIORITAS TINGGI — cocokkan dosis/konsentrasi dgn label resmi; cek APD & keselamatan")

# ------------------------------------------------- 5) VARIETAS
rows = []
for s in seeds:
    if "seed_id" not in s: continue
    rows.append([s.get("variety_name"), s.get("certification_status"), s.get("price_per_seed_rp"),
                 s.get("price_status"), s.get("max_tbs_ton_ha_year"), s.get("cap_status"),
                 s.get("start_harvest_age_year"), s.get("rendemen_pct"), s.get("economic_lifespan_year"),
                 s.get("traits", {}).get("ganoderma_tolerance", ""), s.get("trait_status"),
                 s.get("source_url", "")])
make_sheet("Varietas_Kecambah",
           ["Varietas", "Sertifikasi", "Harga/butir (Rp)", "Status harga", "Plafon (t/ha/th)", "Status plafon",
            "Umur panen", "Rendemen %", "Umur ekonomis", "Toleransi Ganoderma", "Status ciri", "Sumber"],
           rows, [18, 12, 14, 12, 13, 12, 10, 10, 12, 16, 14, 34],
           note="PRIORITAS SEDANG — konfirmasi harga & klaim varietas ke produsen resmi")

# ------------------------------------------------- 6) PARAMETER PLACEHOLDER
params = load("params.json")
prows = []
prows.append(["Harga TBS", "tbs_price_rp_per_kg", params.get("tbs_price_rp_per_kg"), "estimasi",
              "Harga pasar TBS terkini per wilayah"])
for k, v in params.get("fertilizer_price_rp_per_kg", {}).items():
    prows.append(["Harga pupuk (Rp/kg)", f"fertilizer_price_rp_per_kg.{k}", v, "PLACEHOLDER (tak ada di workbook)",
                  "Harga pupuk aktual — T_Harga hanya berisi pestisida"])
for k, v in params.get("pest_treatment_cost_rp_per_ha", {}).items():
    prows.append(["Biaya obat hama (Rp/ha)", f"pest_treatment_cost_rp_per_ha.{k}", v, "PLACEHOLDER",
                  "Biaya aktual per metode (dosis hama masih needs_review)"])
for k in ["pest_level_growth_per_day", "pest_damage_per_day", "pest_base_effectiveness",
          "pest_inspection_cost_rp", "weed_growth_per_day", "spray_base_effectiveness"]:
    prows.append(["Parameter mekanik", k, params.get(k), "gameplay", "Setel lewat uji keseimbangan"])
make_sheet("Parameter_Placeholder",
           ["Kelompok", "Kunci (params.json)", "Nilai sekarang", "Status", "Yang perlu diisi/diverifikasi"],
           prows, [24, 36, 16, 26, 44],
           note="SEDANG — ganti placeholder dgn data nyata; parameter gameplay disetel via uji keseimbangan")

# ------------------------------------------------- 7) KATALOG HAMA GAMEPLAY
rows = []
for p in pcat:
    rows.append([p.get("name"), p.get("type"), p.get("diagnosis_group"), p.get("age_min"), p.get("age_max"),
                 p.get("appear_chance_per_day"), p.get("damage_mult"), p.get("method"),
                 ", ".join(p.get("ingredients", [])) or "(non-kimia)", p.get("curable"), p.get("source_sheet")])
make_sheet("Katalog_Hama_Gameplay",
           ["Hama/penyakit", "Tipe", "Grup diagnosis", "Umur min", "Umur max", "Peluang/hari", "Pengali rusak",
            "Metode", "Bahan aktif", "Dapat sembuh", "Sheet asal"],
           rows, [26, 10, 16, 9, 9, 12, 12, 14, 28, 12, 12],
           note="RENDAH — angka penyeimbang; pastikan metode/bahan cocok kondisi nyata")

wb.save(OUT)
print("Tersimpan:", os.path.abspath(OUT))
print("Sheet:", wb.sheetnames)
