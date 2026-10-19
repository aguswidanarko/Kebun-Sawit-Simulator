#!/usr/bin/env python3
"""
Konversi workbook -> data/*.json runtime (lanjutan import_workbook.py).

Tabel T_* (datar & sudah ternormalisasi oleh penulis workbook) diimpor PENUH
dengan status 'from_workbook'. Sheet hama (Ulat Api/Hama Lain/Bibitan) bermerged
cells & satuan campuran -> diimpor KONSERVATIF dengan nilai mentah dipertahankan
dan status 'needs_review' (tidak dipakai rumus otomatis; lihat GDD §7).

Pemakaian: python3 tools/convert_tables.py
"""
import json, os, re

HERE = os.path.dirname(__file__)
XLSX = os.path.join(HERE, "..", "..", "Std_Produksi_Pupuk_Herbisida_v3.xlsx")
DATA = os.path.join(HERE, "..", "data")

import openpyxl
wb = openpyxl.load_workbook(XLSX, data_only=True)

def write(name, rows, note):
    rows = [{"_catatan": note, "source_version": "v3"}] + rows
    path = os.path.join(DATA, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print("  %-26s %4d baris -> %s" % (name, len(rows) - 1, name))

def num(v):
    if isinstance(v, (int, float)):
        return round(float(v), 3)
    return v

def header_row(ws, key):
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and str(row[0]).strip().lower() == key:
            return r
    return 3

def filled_grid(ws):
    """Kembalikan grid di mana sel merged diisi nilai top-left (untuk sheet hama)."""
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    for rng in ws.merged_cells.ranges:
        tl = grid[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                    grid[r - 1][c - 1] = tl
    return grid

# ---------------------------------------------------------------- T_Pupuk
def conv_pupuk():
    ws = wb["T_Pupuk"]
    hr = header_row(ws, "umur")
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
        if row[0] is None:
            continue
        umur, pupuk = row[0], row[1]
        kg = {"S1": row[2], "S2": row[3], "S3": row[4]}
        kgp = {"S1": row[5], "S2": row[6], "S3": row[7]}
        for lc in ("S1", "S2", "S3"):
            out.append({
                "plant_age_year": int(umur), "fertilizer_code": str(pupuk), "land_class": lc,
                "kg_ha_year": num(kg[lc]), "kg_per_palm": num(kgp[lc]),
                "source_sheet": "T_Pupuk", "source_row": r, "validation_status": "from_workbook"
            })
    return out

# ---------------------------------------------------------------- T_Semprot
def conv_semprot():
    ws = wb["T_Semprot"]
    hr = header_row(ws, "umur_min")
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
        if row[0] is None or row[5] is None:
            continue
        out.append({
            "id": "semprot_%s" % (row[12] if row[12] is not None else r),
            "target_type": "gulma",
            "plant_age_min": int(row[0]), "plant_age_max": int(row[1]),
            "target_location": str(row[2]), "target_organism": str(row[3]),
            "mix_option": row[4],            # pengelompok komponen dalam 1 campuran
            "method": "knapsack",
            "active_ingredient": str(row[5]),
            "dose_value": num(row[6]), "dose_unit": str(row[7]), "per_basis": "ha",
            "water_l_ha_min": num(row[8]), "water_l_ha_max": num(row[9]),
            "equipment": str(row[10]), "concentration_pct": num(row[11]),
            "source_sheet": "T_Semprot", "source_row": num(row[12]),
            "validation_status": "from_workbook"
        })
    return out

# ---------------------------------------------------------------- T_Harga
def conv_harga():
    ws = wb["T_Harga"]
    hr = header_row(ws, "bahan_aktif")
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
        if row[0] is None:
            continue
        src = str(row[4] or "")
        status = "estimasi" if "estimas" in src.lower() else "aktual"
        out.append({
            "product_id": "prod_%s" % (row[5] if row[5] is not None else r),
            "active_ingredient": str(row[0]), "trade_name": str(row[1]),
            "unit": str(row[2]), "price": num(row[3]), "price_year": 2026,
            "price_status": status, "price_source": src,
            "source_sheet": "T_Harga", "source_row": num(row[5]),
            "validation_status": "from_workbook"
        })
    return out

# ---------------------------------------------------------------- Hama (konservatif)
HEADER_TOKENS = {"nama produk", "bahan aktif", "alternatif insektisida",
                 "jenis hama/ penyakit", "jenis pestisida", "metode pengendalian"}

def _is_header(*vals):
    for v in vals:
        if v and str(v).strip().lower() in HEADER_TOKENS:
            return True
    return False

def conv_hama():
    out = []
    # Ulat Api & Hama Lain: layout 13 kolom sama.
    for sheet in ("Ulat Api", "Hama Lain"):
        ws = wb[sheet]
        g = filled_grid(ws)
        group = ""
        for r in range(len(g)):
            row = g[r]
            if len(row) < 12:
                continue
            c1 = str(row[1]).strip() if row[1] else ""
            if c1.endswith(":"):
                group = c1.rstrip(":").strip()
            produk, bahan = row[4], row[5]
            if not produk and not bahan:
                continue
            if _is_header(produk, bahan, c1):
                continue
            out.append({
                "target_type": "hama", "diagnosis_group": group,
                "target_organism_raw": c1, "plant_age_raw": row[2], "method": row[3],
                "trade_name": produk, "active_ingredient": bahan,
                "concentration_raw": row[6], "dose_raw": row[7], "volume_raw": row[8],
                "applicator": row[9], "interval_raw": row[10], "notes": row[11],
                "source_sheet": sheet, "source_row": r + 1, "validation_status": "needs_review"
            })
    # Bibitan: layout 10 kolom (tanpa umur/volume).
    ws = wb["Bibitan"]
    g = filled_grid(ws)
    group = ""
    for r in range(len(g)):
        row = g[r]
        if len(row) < 10:
            continue
        c1 = str(row[1]).strip() if row[1] else ""
        if c1.endswith(":") or c1 in ("HAMA", "PENYAKIT"):
            group = c1.rstrip(":").strip()
        produk, bahan = row[3], row[4]
        if not produk and not bahan:
            continue
        if _is_header(produk, bahan, c1):
            continue
        out.append({
            "target_type": "bibitan", "diagnosis_group": group,
            "target_organism_raw": c1, "method": row[2],
            "trade_name": produk, "active_ingredient": bahan,
            "concentration_raw": row[5], "dose_raw": row[6],
            "applicator": row[7], "interval_raw": row[8], "notes": row[9],
            "source_sheet": "Bibitan", "source_row": r + 1, "validation_status": "needs_review"
        })
    return out

print("Konversi tabel workbook -> data/*.json:")
write("fertilizer_standard.json", conv_pupuk(),
      "Diimpor dari T_Pupuk. kg_ha_year sudah termasuk faktor kelas (S1 basis, S2 +5%, S3 +10%) — JANGAN tambah faktor lagi.")
write("control_standard.json", conv_semprot(),
      "Diimpor dari T_Semprot (norma semprot gulma knapsack). mix_option = pengelompok komponen dalam 1 campuran. Bahan aktif != nama dagang (lihat product_catalog.json).")
write("product_catalog.json", conv_harga(),
      "Diimpor dari T_Harga. price_status aktual/estimasi dipertahankan. Verifikasi label & registrasi sebelum materi edukasi publik.")
write("pest_control.json", conv_hama(),
      "Diimpor KONSERVATIF dari Ulat Api/Hama Lain/Bibitan (merged cells, satuan campuran). Nilai MENTAH dipertahankan; validation_status=needs_review; JANGAN pakai di rumus otomatis sebelum ditinjau agronom.")
print("Selesai.")
