#!/usr/bin/env python3
"""Membuat GDD Induk (dokumen induk) .docx — menambal temuan A & B audit."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join(os.path.dirname(__file__), "..", "..",
                   "GDD_Induk_Kebun_Sawit_Simulator_v0.1.docx")

d = Document()

def h(text, level=1):
    d.add_heading(text, level=level)

def p(text, bold=False):
    par = d.add_paragraph()
    r = par.add_run(text); r.bold = bold
    return par

def bullet(text):
    d.add_paragraph(text, style="List Bullet")

def num(text):
    d.add_paragraph(text, style="List Number")

# ------- Judul -------
t = d.add_heading("KEBUN SAWIT SIMULATOR", level=0)
sub = d.add_paragraph("Dokumen Induk (Master GDD) — v0.1")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("Dokumen ini mengunci keputusan fondasi dan menautkan lampiran: "
                "(1) Database Norma v0.3 dan (2) Skenario Awal & Profil Pemain v0.1. "
                "Dibuat menindaklanjuti audit (temuan A: GDD induk belum ada; B: engine belum dikunci).")

# ------- 1. Visi -------
h("1. Visi & pilar", 1)
p("Simulator kebun sawit yang edukatif dan realistis: pemain membangun kebun dari "
  "nol, mengambil keputusan budidaya berbasis data norma nyata, dan mengelola "
  "ekonomi usaha. Keunggulan utama = kedalaman agronomi yang jujur (bukan sekadar "
  "farm game kasual).")
for b in ["Realistis tapi dapat diseimbangkan: semua angka adalah parameter, bukan klaim lapangan.",
          "Edukatif: diagnosis sebelum tindakan; tidak ada 'satu tombol basmi semua'.",
          "Adil: tanpa pay-to-win; monetisasi kosmetik.",
          "Offline-first: simulasi inti jalan tanpa internet."]:
    bullet(b)

# ------- 2. Keputusan yang dikunci -------
h("2. Keputusan fondasi yang DIKUNCI", 1)
tbl = d.add_table(rows=1, cols=2); tbl.style = "Light Grid Accent 1"
tbl.rows[0].cells[0].text = "Keputusan"
tbl.rows[0].cells[1].text = "Nilai (v0.1)"
rows = [
    ("Engine", "Godot 4.7.2 (GDScript, bukan .NET). Validasi: proyek import bersih."),
    ("Platform target", "Android (HP kelas menengah ke bawah). Renderer gl_compatibility."),
    ("Mode", "Single-player offline dulu. Multiplayer = fase lanjut, tidak wajib online."),
    ("Model waktu", "1 tick = 1 HARI. Rotasi panen 15 hari, pupuk 120 hari, semprot 90 hari (configurable)."),
    ("Modal awal", "Rp100.000.000, sama untuk semua profil (asal cerita berbeda)."),
    ("Harga TBS awal", "Rp2.500/kg (estimasi) + faktor rendemen & kualitas."),
    ("Arsitektur data", "Data-driven: norma di data/*.json (impor terkurasi dari workbook). Excel tidak dibaca saat game jalan."),
]
for k, v in rows:
    c = tbl.add_row().cells; c[0].text = k; c[1].text = v

# ------- 3. Core loop -------
h("3. Core gameplay loop", 1)
for i, s in enumerate([
    "Pilih profil & lahan awal (kebun produktif 1 ha / lahan bersih / lahan berhutan).",
    "Inspeksi zona: game baca umur tanaman, kelas lahan, gulma, hama, cuaca.",
    "Mesin aturan menyaring standar (norma) yang cocok + estimasi bahan, air, HK, biaya.",
    "Pemain memutuskan: kerja sendiri / karyawan / pekerja lepas; cek stok, dana, alat.",
    "Eksekusi: stok berkurang sesuai pemakaian aktual; HK & biaya dicatat.",
    "Efektivitas dihitung (ketepatan diagnosis, metode, cakupan, cuaca, kalibrasi).",
    "Hasil masuk log blok -> memengaruhi gulma/serangan, kesehatan, biaya, produksi mendatang.",
    "Panen tiap rotasi -> jual TBS -> kas -> ekspansi/replanting.",
]):
    num(s)

# ------- 4. Ekonomi -------
h("4. Model ekonomi (rumus inti)", 1)
bullet("Produksi/rotasi = min(ton_ha_tahun_tabel, plafon_varietas) x 1000 x (hari_rotasi/365) x luas_ha.")
bullet("Harga efektif = harga_dasar x faktor_rendemen x kualitas(health). "
       "faktor_rendemen = 1 + (rendemen% - baseline 22%) x 0,01 x sensitivitas.")
bullet("Golden Tree bonus = jumlah_golden x produksi_per_pokok_normal x (multiplier-1), "
       "DITAMBAHKAN di atas plafon (lihat bagian 5).")
bullet("Biaya = bahan + HK aplikasi (terpisah). Norma HK yang tidak ada di workbook "
       "ditandai parameter kosong, TIDAK dikarang.")

# ------- 5. Golden Tree -------
h("5. Golden Tree — keputusan final (temuan D & E)", 1)
p("Masalah audit: jika bonus 5x di-clamp ke plafon blok, Golden Tree jadi tanpa efek. "
  "KEPUTUSAN:", bold=True)
bullet("Produksi NORMAL blok tetap dibatasi plafon varietas.")
bullet("Bonus Golden Tree DITAMBAHKAN di atas plafon (golden_bonus_exceeds_variety_cap=true). "
       "Dengan ~5% pokok, bonus maksimal ~+20% per blok: terasa namun terkendali.")
bullet("Peluang 5% hanya untuk kecambah bersertifikat; di-roll SEKALI di akhir fase "
       "pembibitan (bukan tiap panen). Faktor perawatan memodifikasi peluang pada jendela itu.")
bullet("Visual tandan emas + label 'Golden Tree'; seed/hasil roll dicatat agar dapat diaudit.")

# ------- 6. Lampiran -------
h("6. Sistem & lampiran (dokumen tertaut)", 1)
bullet("LAMPIRAN A — Database Norma v0.3: norma produksi, pupuk, semprot, hama, kecambah, "
       "Golden Tree, model data. Sumber: Std_Produksi_Pupuk_Herbisida_v3.xlsx.")
bullet("LAMPIRAN B — Skenario Awal & Profil Pemain v0.1: 3 profil (Pekerja Aktif, "
       "Anak Pengusaha, Mahasiswa Pertanian), monetisasi kosmetik.")
p("Diferensiasi varietas (temuan F): varietas kini dibedakan pada umur mulai panen, "
  "rendemen, umur ekonomis, toleransi Ganoderma, dan adaptasi lahan — bukan hanya plafon.")

# ------- 7. Scope & roadmap -------
h("7. Scope & roadmap bertahap", 1)
phases = [
    ("Fase 1 — MVP ekonomi (SEKARANG)", "Offline, 1 ha, loop beli->tanam->waktu->panen->jual. "
     "Data-driven. SUDAH ADA scaffold Godot + data produksi asli terimpor."),
    ("Fase 2 — Budidaya", "Pupuk, gulma, hama/penyakit dengan diagnosis; health & efektivitas; cuaca."),
    ("Fase 3 — Profil & UI", "Layar pemilihan profil, tutorial adaptif, misi 7 hari, save/load lengkap."),
    ("Fase 4 — Tenaga kerja & pasar", "Pekerja lepas/karyawan, HK, pasar TBS (pengepul vs pabrik, grading)."),
    ("Fase 5 — Visual 3D", "Kebun 3D, karakter, kendaraan — setelah loop ekonomi solid."),
    ("Fase 6 — Multiplayer & monetisasi", "Online opsional, kosmetik. Tidak mewajibkan koneksi."),
]
for k, v in phases:
    p(k, bold=True); d.add_paragraph(v)

# ------- 8. Status data -------
h("8. Status data & validasi (temuan J)", 1)
bullet("production_standard.json: TERIMPOR dari T_Produksi (123 baris, umur 2-42). from_workbook.")
bullet("fertilizer_standard.json: TERIMPOR dari T_Pupuk (1008 baris; faktor S1/S2/S3 sudah tertanam).")
bullet("control_standard.json: TERIMPOR dari T_Semprot (136 baris gulma; multi-bahan sudah ternormalisasi, mix_option mengelompokkan campuran).")
bullet("product_catalog.json: TERIMPOR dari T_Harga (76 produk, 28 bahan aktif; status aktual/estimasi dipertahankan).")
bullet("pest_control.json: diimpor KONSERVATIF dari Ulat Api/Hama Lain/Bibitan (111 baris, needs_review — merged cells & satuan campuran; BELUM dipakai rumus otomatis).")
bullet("seed_catalog: angka plafon/harga/rendemen = gameplay/estimasi, ditandai perlu verifikasi.")
bullet("WAJIB: cek status registrasi & label produk sebelum dosis dipakai sebagai materi edukasi publik.")
bullet("Tooling: import_workbook.py (audit->data_dictionary.json), convert_tables.py (konversi->data/*.json).")

# ------- 9. Changelog -------
h("9. Changelog", 1)
bullet("v0.1 — Dokumen induk dibuat. Keputusan fondasi dikunci. Golden Tree & model waktu "
       "difinalkan. Data produksi diimpor dari workbook. Proyek Godot scaffold tervalidasi.")

d.save(OUT)
print("Tersimpan:", os.path.abspath(OUT))
